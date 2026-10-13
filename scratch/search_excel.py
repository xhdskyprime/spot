import os
import glob
import openpyxl

def search_in_xlsx(folder):
    files = glob.glob(os.path.join(folder, "**/*.xlsx"), recursive=True) + glob.glob(os.path.join(folder, "**/*.xls"), recursive=True)
    for f in files:
        if f.endswith('.xlsx'):
            try:
                wb = openpyxl.load_workbook(f, read_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    for row in ws.iter_rows(values_only=True):
                        for cell in row:
                            if cell and "lutfi" in str(cell).lower():
                                print(f"Found in {f} (sheet {sheet}): {cell}")
                                break
            except Exception as e:
                pass
        elif f.endswith('.xls'):
            # It's actually HTML in our case, but let's check
            try:
                with open(f, 'r', encoding='utf-8') as file:
                    content = file.read()
                    if "lutfi" in content.lower():
                        print(f"Found in {f} (HTML/xls)")
            except:
                pass

search_in_xlsx("d:\\SPOT")
